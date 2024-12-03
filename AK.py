# Required Libraries
import PyPDF2
import re
import docx
import spacy
import json
import requests
from nltk.corpus import stopwords
from nltk.tokenize import word_tokenize
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from fpdf import FPDF  # For PDF report generation
from transformers import pipeline  # For summarization
from sklearn.feature_extraction.text import CountVectorizer
from language_tool_python import LanguageTool
from googletrans import Translator

# Ensure NLTK dependencies are available
import nltk
nltk.download('punkt')
nltk.download('stopwords')


# Function to Extract Text from PDF
def extract_text(file_path):
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
        return text
    except Exception as e:
        print(f"Error extracting text: {e}")
        return ""


# Function to Preprocess Text
def preprocess_text(text):
    stop_words = set(stopwords.words('english'))
    text = re.sub(r'\s+', ' ', text)  # Remove extra spaces
    text = re.sub(r'[^\w\s]', '', text)  # Remove punctuation
    tokens = word_tokenize(text.lower())
    filtered_tokens = [word for word in tokens if word not in stop_words]
    return " ".join(filtered_tokens)


# Function to Extract Skills
def extract_skills(text, skill_set):
    extracted_skills = [skill for skill in skill_set if skill.lower() in text.lower()]
    return extracted_skills


# Function to Match Resume to Job Description
def match_resume_to_job(resume_text, job_description):
    vectorizer = TfidfVectorizer()
    vectors = vectorizer.fit_transform([resume_text, job_description])
    similarity = cosine_similarity(vectors[0:1], vectors[1:2])
    return similarity[0][0]


# Function to Rank Resumes
def rank_resumes(resume_texts, job_description):
    rankings = []
    for resume_text in resume_texts:
        score = match_resume_to_job(resume_text, job_description)
        rankings.append((resume_text, score))
    return sorted(rankings, key=lambda x: x[1], reverse=True)


# Function to Load Skill Set
def load_skill_set(skill_file):
    try:
        with open(skill_file, 'r') as file:
            skills = file.read().splitlines()
        return set(skills)
    except Exception as e:
        print(f"Error loading skill set: {e}")
        return set()


# Function to Generate Summaries
def generate_summary(text, max_length=150):
    try:
        summarizer = pipeline("summarization")
        summary = summarizer(text, max_length=max_length, min_length=30, do_sample=False)
        return summary[0]['summary_text']
    except Exception as e:
        print(f"Error generating summary: {e}")
        return "Unable to summarize text."


# Function to Identify Skill Gaps
def identify_skill_gaps(extracted_skills, required_skills):
    return list(required_skills - set(extracted_skills))


# Function to Generate a PDF Report
def generate_pdf_report(rankings, skill_gaps, output_path="resume_report.pdf"):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, txt="Resume Screening Report", ln=True, align="C")
    pdf.ln(10)

    # Rankings
    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(0, 10, txt="Resume Rankings:", ln=True)
    pdf.set_font("Arial", size=12)
    for idx, (resume, score) in enumerate(rankings, 1):
        pdf.cell(0, 10, txt=f"{idx}. Resume Score: {score:.2f}", ln=True)

    pdf.ln(10)

    # Skill Gaps
    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(0, 10, txt="Skill Gaps:", ln=True)
    pdf.set_font("Arial", size=12)
    for gap in skill_gaps:
        pdf.cell(0, 10, txt=f"- {gap}", ln=True)

    pdf.output(output_path)
    print(f"PDF report generated at {output_path}")
    

def extract_keywords(text, top_n=10):
    try:
        vectorizer = CountVectorizer(stop_words='english', max_features=top_n)
        term_matrix = vectorizer.fit_transform([text])
        keywords = vectorizer.get_feature_names_out()
        return keywords
    except Exception as e:
        print(f"Error extracting keywords: {e}")
        return []

# Load a pretrained SpaCy model
nlp = spacy.load("en_core_web_sm")

def extract_entities(text):
    try:
        doc = nlp(text)
        entities = {ent.label_: [] for ent in doc.ents}
        for ent in doc.ents:
            entities[ent.label_].append(ent.text)
        return entities
    except Exception as e:
        print(f"Error extracting entities: {e}")
        return {}
        
def find_missing_skills_in_job(extracted_skills, job_description_skills):
    return [skill for skill in extracted_skills if skill not in job_description_skills]

def parse_resume_sections(text):
    sections = {
        "Education": [],
        "Experience": [],
        "Skills": [],
        "Others": []
    }
    current_section = "Others"
    for line in text.split('\n'):
        if "education" in line.lower():
            current_section = "Education"
        elif "experience" in line.lower():
            current_section = "Experience"
        elif "skills" in line.lower():
            current_section = "Skills"
        else:
            sections[current_section].append(line)
    return sections
    
def calculate_job_suitability(resume_text, job_descriptions):
    scores = {}
    for idx, job in enumerate(job_descriptions, 1):
        score = match_resume_to_job(resume_text, preprocess_text(job))
        scores[f"Job {idx}"] = score
    return scores
    

def generate_summary_report(ranked_resumes, skill_gaps, output_path="summary_report.json"):
    report = {
        "ranked_resumes": [{"score": score, "resume": text[:100] + "..."} for text, score in ranked_resumes],
        "skill_gaps": skill_gaps
    }
    with open(output_path, 'w') as file:
        json.dump(report, file, indent=4)
    print(f"Summary report saved to {output_path}")
    

def fetch_job_description_from_api(api_url, api_key, job_id):
    headers = {"Authorization": f"Bearer {api_key}"}
    response = requests.get(f"{api_url}/jobs/{job_id}", headers=headers)
    if response.status_code == 200:
        return response.json().get("job_description", "")
    else:
        print(f"Failed to fetch job description: {response.status_code}")
        return ""

def recommend_candidates(candidate_data, job_description):
    vectorizer = TfidfVectorizer()
    vectors = vectorizer.fit_transform(candidate_data + [job_description])
    similarities = cosine_similarity(vectors[-1], vectors[:-1])
    ranked_candidates = sorted(
        enumerate(similarities[0]), key=lambda x: x[1], reverse=True
    )
    return ranked_candidates
    
def collect_feedback(resume_id, feedback_text, feedback_log="feedback.txt"):
    try:
        with open(feedback_log, "a") as log:
            log.write(f"Resume ID: {resume_id}, Feedback: {feedback_text}\n")
        print("Feedback recorded successfully.")
    except Exception as e:
        print(f"Error recording feedback: {e}")

def detect_overqualification(resume_text, job_description):
    overqualified_keywords = ["senior", "lead", "manager", "director", "advanced"]
    resume_keywords = [word for word in overqualified_keywords if word in resume_text.lower()]
    job_keywords = [word for word in overqualified_keywords if word in job_description.lower()]
    return len(resume_keywords) > len(job_keywords)

def check_resume_formatting(text):
    issues = []
    if not any(header in text.lower() for header in ["education", "experience", "skills"]):
        issues.append("Missing key sections (Education, Experience, Skills).")
    if len(text.splitlines()) < 10:
        issues.append("Resume might be too short.")
    return issues if issues else ["No formatting issues detected."]

def weighted_skill_score(extracted_skills, weighted_skills):
    total_weight = 0
    matched_weight = 0
    for skill, weight in weighted_skills.items():
        total_weight += weight
        if skill.lower() in [s.lower() for s in extracted_skills]:
            matched_weight += weight
    return matched_weight / total_weight if total_weight > 0 else 0

def detect_career_gaps(text):
    import re
    years = re.findall(r'\b(19|20)\d{2}\b', text)
    years = sorted(map(int, years))
    gaps = [(years[i], years[i+1]) for i in range(len(years) - 1) if years[i+1] - years[i]]

# Main Script
if __name__ == "__main__":
    # Define paths
    resumes = ["resume1.pdf", "resume2.pdf"]  # Add paths to your resume files
    job_description_path = "job_description.txt"
    skill_file_path = "skills.txt"
    feedback_log = "feedback.txt"  # Log for feedback collection

    # Load job description and skills
    try:
        with open(job_description_path, 'r') as jd_file:
            job_description = jd_file.read()
    except FileNotFoundError:
        print("Job description file not found.")
        job_description = ""

    skill_set = load_skill_set(skill_file_path)

    # Process resumes
    processed_resumes = []
    extracted_skills_all = []
    for resume_path in resumes:
        print(f"Processing {resume_path}...")
        resume_text = extract_text(resume_path)
        preprocessed_text = preprocess_text(resume_text)
        skills = extract_skills(preprocessed_text, skill_set)
        extracted_skills_all.append(skills)
        print(f"Extracted Skills: {skills}")
        processed_resumes.append(preprocessed_text)

    # Rank resumes
    ranked_resumes = rank_resumes(processed_resumes, preprocess_text(job_description))

    # Identify skill gaps
    required_skills = set(skill_set)
    combined_skills = set(skill for skills in extracted_skills_all for skill in skills)
    skill_gaps = identify_skill_gaps(combined_skills, required_skills)

    # Generate PDF report
    generate_pdf_report(ranked_resumes, skill_gaps)

    print("\nResume Rankings (Higher is Better):")
    for idx, (resume, score) in enumerate(ranked_resumes, 1):
        print(f"{idx}. Resume Score: {score:.2f}")

    print("\nSkill Gaps:")
    for gap in skill_gaps:
        print(f"- {gap}")

    # Collect feedback
    for idx, resume_text in enumerate(resumes, 1):
        feedback = input(f"Enter feedback for Resume {idx}: ")
        collect_feedback(resume_id=idx, feedback_text=feedback, feedback_log=feedback_log)

    # Detect overqualification
    for idx, resume_text in enumerate(resumes, 1):
        is_overqualified = detect_overqualification(resume_text, job_description)
        if is_overqualified:
            print(f"Resume {idx} is overqualified for the job.")

    # Check for career gaps
    for idx, resume_text in enumerate(resumes, 1):
        career_gaps = detect_career_gaps(resume_text)
        if career_gaps:
            print(f"Resume {idx} has career gaps: {career_gaps}")

    # Evaluate writing quality
    for idx, resume_text in enumerate(resumes, 1):
        quality = evaluate_writing_quality(resume_text)
        if quality['issues_found'] > 0:
            print(f"Resume {idx} has writing quality issues: {quality['details']}")
        else:
            print(f"Resume {idx} has no writing quality issues.")

    # Check for red flags
    for idx, resume_text in enumerate(resumes, 1):
        red_flags = check_red_flags(resume_text)
        if red_flags:
            print(f"Resume {idx} has the following red flags: {red_flags}")
        else:
            print(f"Resume {idx} has no red flags.")
    
    # Check resume formatting
    for idx, resume_text in enumerate(resumes, 1):
        formatting_issues = check_resume_formatting(resume_text)
        if formatting_issues:
            print(f"Resume {idx} has formatting issues: {formatting_issues}")
        else:
            print(f"Resume {idx} has no formatting issues.")

    # Detect duplicates
    duplicates = check_duplicates(resumes)
    if duplicates:
        for duplicate in duplicates:
            print(duplicate)
    else:
        print("No duplicate resumes detected.")
    
    # Extract contact info
    for idx, resume_text in enumerate(resumes, 1):
        contact_info = extract_contact_info(resume_text)
        print(f"Resume {idx} contact info: {contact_info}")

    # Generate a summary report (you can specify the file path)
    generate_summary_report(ranked_resumes, skill_gaps, output_path="summary_report.json")
